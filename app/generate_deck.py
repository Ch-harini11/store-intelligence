import os
import sys

def main():
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
    except ImportError:
        print("[PPTX ERROR] python-pptx is not installed. To generate PowerPoint file directly, run: pip install python-pptx")
        return

    print("Generating PowerPoint slides...")
    prs = Presentation()
    
    # Slide dimensions (16:9 widescreen)
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Colors
    c_bg = RGBColor(9, 10, 15)       # Deep dark
    c_white = RGBColor(255, 255, 255)
    c_text = RGBColor(148, 163, 184)  # Slate secondary
    c_purple = RGBColor(157, 78, 221) # Purple accent
    c_emerald = RGBColor(6, 214, 160) # Emerald accent
    c_blue = RGBColor(58, 134, 200)   # Blue accent

    def set_dark_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = c_bg

    def add_title(slide, text, color=c_white):
        tx_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.5), Inches(11.83), Inches(1))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = "Arial"
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = color

    # Slide 1: Title
    slide1 = prs.slides.add_slide(blank_layout)
    set_dark_background(slide1)
    
    tb1 = slide1.shapes.add_textbox(Inches(0.75), Inches(2.2), Inches(11.83), Inches(3))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p1_title = tf1.paragraphs[0]
    p1_title.text = "Purplle Store Intelligence System"
    p1_title.font.name = "Arial"
    p1_title.font.size = Pt(48)
    p1_title.font.bold = True
    p1_title.font.color.rgb = c_purple
    
    p1_sub = tf1.add_paragraph()
    p1_sub.text = "Bringing E-Commerce Analytics Parity to Physical Retail Stores"
    p1_sub.font.name = "Arial"
    p1_sub.font.size = Pt(20)
    p1_sub.font.color.rgb = c_white
    p1_sub.space_before = Pt(20)

    p1_pres = tf1.add_paragraph()
    p1_pres.text = "Apex Retail Analytics Team | Hackathon Submission"
    p1_pres.font.name = "Arial"
    p1_pres.font.size = Pt(14)
    p1_pres.font.color.rgb = c_text
    p1_pres.space_before = Pt(30)

    # Content Slides Template
    slides_data = [
        {
            "title": "The Offline Retail Data Blind Spot",
            "points": [
                "Online channels track every click, bounce, session duration, and cart drop-off in real-time.",
                "Physical stores represent the vast majority of revenue but remain complete data blind spots.",
                "Store managers have no structured metrics on customer pathways, shelf interaction, or line abandonments.",
                "Apex Retail operates 40 physical stores across 8 cities in statistical darkness before this system."
            ],
            "accent": c_purple
        },
        {
            "title": "The Solution: End-to-End Store Intelligence",
            "points": [
                "Edge Video Pipeline: Converts existing CCTV surveillance streams into structured event logs.",
                "Zero Human Overhead: Fully automated YOLOv8 person tracking and zone mapping.",
                "Clean Edge Operations: Local lightweight SQLite storage mapping shopper sessions with POS correlation.",
                "Actionable Dashboard Surface: Renders conversion funnels, occupancy, peak hours, and warnings."
            ],
            "accent": c_emerald
        },
        {
            "title": "Technical Pipeline Architecture",
            "points": [
                "CCTV Video Streams: 1080p feeds processed per frame (Entrance, Floor, Billing).",
                "YOLOv8 + ByteTrack: High-speed person detection and persistent visitor ID tracking.",
                "Spatial Zone OpenCV: 2D polygon intersection checks coordinates and calculates zone dwell times.",
                "SQLite DB + FastAPI: Ingests batches with bulk deduplication and calculates shop funnels.",
                "Glassmorphic UI: Dynamic HTML/CSS/JS dashboard auto-polls active metrics every 3 seconds."
            ],
            "accent": c_blue
        },
        {
            "title": "Spatial Zone Intelligence",
            "points": [
                "Modular Store Geometry: Store layout mapped into 2D pixel coordinate polygons.",
                "Containment Engine: Bottom-center bounding box points checked via OpenCV pointPolygonTest.",
                "Visitor State Tracking: Chronologically logs ZONE_ENTER, ZONE_EXIT, and ZONE_DWELL transitions.",
                "Audit Path Trails: Session sequence variables log exact visitor browsing pathway chronologies."
            ],
            "accent": c_purple
        },
        {
            "title": "Operational Anomalies Rules Engine",
            "points": [
                "CAMERA_FAILURE (Critical): Alerts if camera ingestion goes quiet for over 2 minutes.",
                "LONG_QUEUE (Warning): Warns if the number of active shoppers waiting at checkout exceeds 5.",
                "FOOTFALL_SPIKE (Warning): Flags if shopper entry frequency rises above 3x the running hourly average.",
                "EMPTY_STORE (Warning): Alerts managers if occupancy is 0 for 30+ minutes during business hours."
            ],
            "accent": c_emerald
        },
        {
            "title": "Conversion Funnel & Analytics Metrics",
            "points": [
                "Standard Retail Funnel: Unique Visitors Entered -> Visited Shelves -> Checked Out at Billing.",
                "Conversion Rate Calculation: Correlates visitors in billing zone in 5-minute windows with POS data.",
                "Sessionization Logic: Prevents customer re-entry from double-counting unique visitor footfalls.",
                "Peak Hour Detection: Isolates busy traffic slots to optimize local employee shifts."
            ],
            "accent": c_blue
        },
        {
            "title": "Premium Glassmorphic UI Dashboard",
            "points": [
                "Premium Visual Aesthetic: Designed with glassmorphism blur filters and dark neon styling.",
                "Auto-Polling Framework: Dynamic JS requests metrics and occupancies every 3 seconds.",
                "Floor Occupancy Grid: CSS grid displays live customer density in zones visually.",
                "Control Panel: Integrates DB resets and mock seeding triggers for interactive reviews."
            ],
            "accent": c_purple
        },
        {
            "title": "Production Readiness & Containerization",
            "points": [
                "Single Command Boot: 'docker-compose up --build -d' compiles everything in under 5 minutes.",
                "Persistent Local Volume: Mounts host SQLite database file to preserve metrics history.",
                "Bulk Ingest Optimization: Pre-fetches IDs in a single operation, keeping endpoints idempotent.",
                "100% Passing Test Coverage: Pytest validates metrics, funnels, and anomaly rules successfully."
            ],
            "accent": c_emerald
        },
        {
            "title": "Conclusion: Edge-Ready Retail Scale",
            "points": [
                "Edge-Native Deployment: Built to run locally on low-cost edge computers inside stores.",
                "Highly Modular Codebase: Clean independent routers for metrics, funnels, health, and anomalies.",
                "Centralized Scale-Ready: SQLAlchemy ORM allows switching database engines with a single URL string.",
                "Compliance: Full coverage for all 5 Hackathon Acceptance Gates."
            ],
            "accent": c_blue
        }
    ]

    for sd in slides_data:
        slide = prs.slides.add_slide(blank_layout)
        set_dark_background(slide)
        add_title(slide, sd["title"], sd["accent"])
        
        tb = slide.shapes.add_textbox(Inches(0.75), Inches(1.8), Inches(11.83), Inches(5))
        tf = tb.text_frame
        tf.word_wrap = True
        
        for idx, pt in enumerate(sd["points"]):
            if idx == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            
            p.text = "• " + pt
            p.font.name = "Arial"
            p.font.size = Pt(18)
            p.font.color.rgb = c_white
            p.space_after = Pt(20)

    os.makedirs("docs", exist_ok=True)
    out_path = os.path.join("docs", "presentation.pptx")
    prs.save(out_path)
    print(f"Presentation successfully compiled and saved to {out_path}!")

if __name__ == "__main__":
    main()
